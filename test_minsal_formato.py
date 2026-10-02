import ast
from copy import deepcopy
from datetime import date
from io import BytesIO
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import Mock

from docx import Document
from minsal_formato import completeness, generate_docx, format_tables, rows_from_editor
import pandas as pd


def fixtures():
    est = {}
    reports = []
    for eid, level, included in [("a", "rojo", "Sí"), ("b", "rojo", "Sí"), ("c", "amarillo", "Sí"), ("d", "rojo", "No"), ("v", "verde", "No")]:
        e = {"nombre": f"Establecimiento {eid}", "nivel": level, "codigo_deis": eid,
             "rut": "Referencia de prueba", "denominador": 1000000, "numerador": 200000,
             "pct_2026": 20, "pct_2025": 15, "brecha": 4, "variacion": 5}
        est[eid] = e
        f = {"incluido_minsal": included, "riesgo_minsal": level.capitalize(), "pct_minsal": "20",
             "observacion_cotejo": "", "por_informar": "No existen", "avance_anteriores": "No existen compromisos anteriores.",
             "acciones": [{"causa": "Necesidad de abastecimiento", "medida": "Revisión del plan de compras", "compromiso": "Iniciar proceso competitivo", "responsable": "Referente de prueba", "plazo": "2026-11-20", "avance": "En ejecución"}],
             "mecanismos": ["Licitación"], "detalle_mecanismo": "Proceso competitivo de suministros", "fecha_reunion": "2026-10-20",
             "mensual": [{"mes": f"Mes {i}", "fecha": f"2026-{9+i:02d}-20", "estado": "Programado", "observacion": "Control programado"} for i in range(1, 4)],
             "revision_recurrentes": "Sí", "detalle_recurrentes": "Revisión mensual de suministros recurrentes",
             "seguimiento_reforzado": "Sí", "modalidad_seguimiento": "Reunión de seguimiento", "fecha_seguimiento": "2026-10-25",
             "respaldos": [{"oc": "OC-PRUEBA", "descripcion": "Detalle de orden de compra", "url": "https://example.org/respaldo"}]}
        reports.append({"establecimiento_id": eid, "establecimiento_nombre": e["nombre"], "year": 2026, "reporte_id": "R3", "estado": "enviado", "formato_tipo": f})
    return est, reports


class FormatoTests(unittest.TestCase):
    def test_requirements_for_both_risks(self):
        est, reports = fixtures()
        for r in reports[:-1]:
            self.assertEqual(completeness(r, est[r["establecimiento_id"]]), [])
        r = deepcopy(reports[0]); r["formato_tipo"]["acciones"][0]["responsable"] = ""
        self.assertTrue(any("responsable" in x for x in completeness(r, est["a"])))
        r = deepcopy(reports[0]); r["formato_tipo"]["mensual"][1]["fecha"] = ""
        self.assertTrue(any("control mensual" in x for x in completeness(r, est["a"])))

    def test_por_informar_and_discrepancies_block_complete_status(self):
        est, reports = fixtures()
        r = deepcopy(reports[0]); r["formato_tipo"]["por_informar"] = "Pendiente de revisión"
        self.assertTrue(any("reclasifique" in x for x in completeness(r, est["a"])))
        r = deepcopy(reports[0]); r["formato_tipo"]["pct_minsal"] = "23,5"
        self.assertTrue(any("Fundamente" in x for x in completeness(r, est["a"])))
        r["formato_tipo"]["observacion_cotejo"] = "Diferencia documentada con detalle OC"
        self.assertEqual(completeness(r, est["a"]), [])

    def test_pending_legacy_and_period_isolation(self):
        est, reports = fixtures()
        reports[0].pop("formato_tipo")
        reports[1]["reporte_id"] = "R2"
        result = format_tables(est, reports, 2026, "R3", {})
        self.assertEqual(len(result[2]), 4)
        self.assertEqual(len(result[4]), 2)
        self.assertNotIn("b", result[1])
        self.assertNotIn("v", result[0])
        self.assertTrue(result[5])

    def test_editor_ignores_empty_rows_and_preserves_unicode(self):
        f = pd.DataFrame([{"A": None, "B": None}, {"A": "Revisión", "B": date(2026, 10, 2)}])
        self.assertEqual(rows_from_editor(f, {"a": "A", "b": "B"}), [{"a": "Revisión", "b": "2026-10-02"}])

    def test_docx_preserves_sections_and_duplicates_blocks(self):
        est, reports = fixtures()
        before = deepcopy(reports)
        meta = {"remitente": "Referente de prueba", "cargo": "Coordinación", "correo": "referente@example.org", "fecha_envio": "2026-11-20", "export_estado": "PARA REVISIÓN Y VISADO INSTITUCIONAL"}
        content, warnings = generate_docx(est, reports, 2026, "R3", {"label": "3° Reporte", "periodo": "Julio–Septiembre 2026"}, meta)
        self.assertEqual(warnings, [])
        self.assertEqual(reports, before)
        doc = Document(BytesIO(content))
        paragraphs = [p.text for p in doc.paragraphs]
        headings = [t for t in paragraphs if t[:3] in [f"{i}. " for i in range(1, 9)]]
        self.assertEqual(len(headings), 8)
        bodies = list(doc.element.body)
        for name, start, end in [("Establecimiento c", "4. ", "5. "), ("Establecimiento a", "5. ", "6. "), ("Establecimiento b", "5. ", "6. "), ("Establecimiento d", "6. ", "7. ")]:
            body_text = ["".join(e.itertext()) for e in bodies]
            a = next(i for i, t in enumerate(body_text) if t.startswith(start))
            b = next(i for i, t in enumerate(body_text) if t.startswith(end))
            self.assertTrue(any(name in t for t in body_text[a:b]), name)
        self.assertEqual(len(doc.tables), 14)  # consolidado, cotejo, four blocks, pendientes
        self.assertTrue(any(p.startswith("Firma:") and "____" in p for p in paragraphs))
        Path("/tmp/simotd_formato_qa.docx").write_bytes(content)

    def test_failed_save_never_updates_local_cache(self):
        source = Path(__file__).with_name("app.py").read_text()
        tree = ast.parse(source)
        func = next(n for n in tree.body if isinstance(n, ast.FunctionDef) and n.name == "_save_json_persistent")
        scope = {"Path": Path, "json": json, "PersistenceError": RuntimeError,
                 "_gh_cfg": Mock(return_value={"repo": "test"}),
                 "_gh_read": Mock(return_value=({}, "oldsha")), "_gh_write": Mock(return_value=False)}
        exec(compile(ast.Module(body=[func], type_ignores=[]), "save_test", "exec"), scope)
        with tempfile.TemporaryDirectory() as d:
            path = Path(d) / "reports.json"; path.write_text('["old"]')
            with self.assertRaises(RuntimeError): scope["_save_json_persistent"](path, path.name, ["new"])
            self.assertEqual(path.read_text(), '["old"]')
            scope["_gh_write"].return_value = True
            scope["_gh_read"].side_effect = [({}, "oldsha"), (["new"], "newsha")]
            scope["_save_json_persistent"](path, path.name, ["new"])
            self.assertEqual(json.loads(path.read_text()), ["new"])


if __name__ == "__main__":
    unittest.main()
