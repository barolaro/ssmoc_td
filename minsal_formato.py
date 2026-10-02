"""Captura y exportación del Formato Tipo MINSAL 2026, compatible con reportes previos."""
from copy import deepcopy
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from urllib.parse import urlparse

import pandas as pd

ACCIONES = {
    "causa": "Principales causas", "medida": "Medidas implementadas",
    "compromiso": "Compromiso", "responsable": "Responsable",
    "plazo": "Plazo comprometido", "avance": "Estado de avance",
}
MESES = {"mes": "Mes de seguimiento", "fecha": "Fecha de control", "estado": "Estado de cumplimiento", "observacion": "Observación"}
RESPALDOS = {"oc": "OC o referencia", "descripcion": "Descripción del respaldo", "url": "Enlace institucional"}
RECLASIFICACIONES = {"oc": "Orden de compra", "mecanismo": "Mecanismo real", "respaldo": "Referencia o enlace del respaldo"}
ANTERIORES = {"periodo": "Reporte de origen", "compromiso": "Compromiso anterior", "avance": "Avance actual", "responsable": "Responsable", "fecha": "Fecha de actualización"}
ESTADOS = ["Programado", "En ejecución", "Cumplido", "Incumplido", "No aplica con justificación"]


def number_chile(value, decimals=2):
    try:
        return f"{float(value):,.{decimals}f}".replace(",", "_").replace(".", ",").replace("_", ".")
    except (TypeError, ValueError):
        return text(value)


FIELD_HELP = {'avance_anteriores': 'Resuma qué se ha cumplido de los reportes anteriores y qué sigue pendiente. Si no existen compromisos previos, indíquelo expresamente.', 'mecanismos': 'Seleccione la alternativa prevista para regularizar el abastecimiento. Puede marcar más de una; explique Otro en el detalle.', 'detalle_mecanismo': 'Indique qué bienes o servicios se regularizarán, mediante qué proceso y cuál es el próximo paso. Ejemplo: iniciar licitación de insumos durante noviembre.', 'fecha_reunion': 'Registre la fecha acordada o programada para revisar el plan con el Servicio. No registre una reunión como realizada si aún está pendiente.', 'revision_recurrentes': 'Indique si revisaron las compras que se repiten por trato directo. Si marca No, explique el motivo y cómo abordarán la revisión.', 'detalle_recurrentes': 'Identifique los bienes o servicios recurrentes, su frecuencia y el resultado de la revisión. Ejemplo: suministro mensual con oportunidad de licitación.', 'seguimiento_reforzado': 'Indique el compromiso de seguimiento adicional para el próximo período. Antes de enviar debe estar definido.', 'modalidad_seguimiento': 'Describa cómo controlarán los compromisos. Ejemplo: reunión mensual del equipo de abastecimiento con revisión del plan.', 'fecha_seguimiento': 'Indique la fecha programada para el seguimiento reforzado.', 'por_informar': 'Revise las OC cuyo mecanismo aparece como Por Informar. Marque No existen solo después de revisar; si fueron reclasificadas, registre el mecanismo real y su respaldo.'}
COLUMN_HELP = {'causa': 'Explique la situación que originó el trato directo y su contexto; evite descripciones generales.', 'medida': 'Describa la acción ya implementada. Si aún no se ha ejecutado, indíquelo y explique la situación actual.', 'compromiso': 'Indique una acción concreta y verificable. Ejemplo: publicar una licitación para el suministro identificado.', 'responsable': 'Identifique a la persona o cargo responsable de ejecutar y dar seguimiento a esta acción.', 'plazo': 'Ingrese la fecha comprometida en formato AAAA-MM-DD. Ejemplo: 2026-11-20.', 'avance': 'Informe el avance real, los resultados y lo que queda pendiente. Ejemplo: bases en elaboración; pendiente revisión técnica.', 'mes': 'Identifique el mes de control. Ejemplo: octubre 2026. Puede agregar meses hasta normalizar el indicador.', 'fecha': 'Ingrese la fecha del control o actualización en formato AAAA-MM-DD. Distinga en el estado si está programado o realizado.', 'estado': 'Seleccione el estado real del compromiso. Programado corresponde a una actividad futura; Cumplido requiere un resultado verificable.', 'observacion': 'Registre el resultado del control, dificultades y próximos pasos. Justifique cualquier estado No aplica.', 'oc': 'Ingrese el identificador de la OC o referencia del documento. Ejemplo: código de orden de compra correspondiente al respaldo.', 'descripcion': 'Explique qué acredita el documento. Ejemplo: detalle de OC que respalda la cifra o cotizaciones consultadas.', 'url': 'Pegue un enlace HTTPS al repositorio institucional. Compruebe que el equipo revisor tenga acceso; no incluya claves.', 'mecanismo': 'Indique el mecanismo de compra real de la OC reclasificada. Debe coincidir con los antecedentes de respaldo.', 'respaldo': 'Identifique el documento o enlace institucional que sustenta la reclasificación de esta OC.', 'periodo': 'Indique el reporte en que se asumió el compromiso. Ejemplo: 2.º reporte 2026.'}

def text(value):
    if value is None:
        return ""
    if isinstance(value, (date, datetime)):
        return value.isoformat()[:10]
    try:
        if pd.isna(value):
            return ""
    except (TypeError, ValueError):
        pass
    return str(value).strip()


def date_value(value):
    try:
        return date.fromisoformat(text(value)[:10])
    except ValueError:
        return None


def fecha_chile(value):
    d = date_value(value)
    return d.strftime("%d-%m-%Y") if d else text(value)


def rows_from_editor(frame, columns):
    return [{k: text(row.get(label)) for k, label in columns.items()}
            for row in frame.to_dict("records")
            if any(text(row.get(label)) for label in columns.values())]


def legacy_action(report):
    return {"causa": text(report.get("causas_desc")), "medida": text(report.get("med_desc")),
            "compromiso": text(report.get("compromisos")),
            "responsable": text(report.get("resp_nombre")), "plazo": text(report.get("fecha_comp")),
            "avance": ""}


def revision(report, establishment):
    """No infiere la clasificación de la comunicación a partir del CSV del Servicio."""
    f = report.get("formato_tipo", {})
    included = f.get("incluido_minsal", "Pendiente de cotejo")
    official = f.get("riesgo_minsal", "Sin cotejar")
    same = "No aplica" if included == "No" else "Pendiente"
    if included == "Sí" and official in ("Rojo", "Amarillo", "Verde"):
        same = "Sí" if official.lower() == establishment.get("nivel") else "No"
    return included, official, same


def completeness(report, establishment):
    f = report.get("formato_tipo", {})
    errors = []
    actions = f.get("acciones", [])
    if not actions:
        errors.append("Registre al menos una acción con causa, medida, compromiso, responsable, plazo y avance.")
    for i, a in enumerate(actions, 1):
        for field, label in ACCIONES.items():
            if not text(a.get(field)):
                errors.append(f"Acción {i}: complete {label.lower()}.")
        if text(a.get("plazo")) and not date_value(a.get("plazo")):
            errors.append(f"Acción {i}: indique el plazo como AAAA-MM-DD.")
    status = f.get("por_informar", "Pendiente de revisión")
    if status == "Pendiente de revisión":
        errors.append("Revise y reclasifique las OC Por Informar antes del envío.")
    if status == "Reclasificadas":
        if not f.get("reclasificaciones"):
            errors.append("Registre las OC reclasificadas y su respaldo.")
        for a in f.get("reclasificaciones", []):
            if not all(text(a.get(k)) for k in RECLASIFICACIONES):
                errors.append("Complete OC, mecanismo real y respaldo en cada reclasificación.")
    for a in f.get("respaldos", []):
        if not text(a.get("descripcion")) or not (text(a.get("oc")) or text(a.get("url"))):
            errors.append("Cada respaldo requiere descripción y OC/referencia o enlace.")
        if text(a.get("url")):
            u = urlparse(text(a["url"]))
            if u.scheme != "https" or not u.netloc or u.username or u.password:
                errors.append("Use enlaces HTTPS sin credenciales para los respaldos.")
    if not text(f.get("avance_anteriores")):
        errors.append("Informe el avance de compromisos anteriores o justifique que no existen.")
    for a in f.get("anteriores", []):
        if not all(text(a.get(k)) for k in ANTERIORES) or not date_value(a.get("fecha")):
            errors.append("Complete todos los campos y la fecha de los compromisos anteriores.")
    if establishment.get("nivel") == "rojo":
        if not date_value(f.get("fecha_reunion")):
            errors.append("Programe la reunión técnica de seguimiento.")
        if not f.get("mecanismos") or not text(f.get("detalle_mecanismo")):
            errors.append("Identifique el mecanismo competitivo y la oportunidad de regularización.")
        monthly = f.get("mensual", [])
        if len(monthly) < 3:
            errors.append("Registre o programe los tres controles mensuales del seguimiento.")
        if len({text(a.get("mes")) for a in monthly}) != len(monthly):
            errors.append("No repita el mes de seguimiento.")
        for a in monthly:
            if not all(text(a.get(k)) for k in MESES) or not date_value(a.get("fecha")):
                errors.append("Complete mes, fecha, estado y observación en cada control mensual.")
            if a.get("estado") not in ESTADOS:
                errors.append("Seleccione un estado válido en cada control mensual.")
    if establishment.get("nivel") == "amarillo":
        if f.get("revision_recurrentes") not in ("Sí", "No") or not text(f.get("detalle_recurrentes")):
            errors.append("Informe la revisión de compras recurrentes y su detalle o justificación.")
        if f.get("seguimiento_reforzado") != "Sí" or not text(f.get("modalidad_seguimiento")) or not date_value(f.get("fecha_seguimiento")):
            errors.append("Comprometa seguimiento reforzado, indicando modalidad y fecha.")
    return list(dict.fromkeys(errors))


def render_complement(report, level, key, year, rid, previous):
    import streamlit as st
    f = deepcopy(report.get("formato_tipo", {}))
    def choice(label, options, field):
        value = f.get(field, options[0])
        return st.selectbox(label, options, index=options.index(value) if value in options else 0, key=f"{key}_{field}", help=FIELD_HELP.get(field))
    def editor(label, field, columns, defaults=None, status_column=None):
        st.markdown(f"**{label}**")
        source = f.get(field, defaults or [])
        frame = pd.DataFrame([{label: text(row.get(k)) for k, label in columns.items()} for row in source], columns=list(columns.values()))
        configs = {label: st.column_config.TextColumn(label, help=COLUMN_HELP.get(k)) for k, label in columns.items()}
        if status_column:
            configs[columns[status_column]] = st.column_config.SelectboxColumn(columns[status_column], options=ESTADOS, help=COLUMN_HELP[status_column])
        edited = st.data_editor(frame, num_rows="dynamic", hide_index=True, use_container_width=True,
                                column_config=configs, key=f"{key}_{field}")
        f[field] = rows_from_editor(edited, columns)

    st.subheader("Plan de acción y estado de avance")
    st.caption("Agregue una fila por acción. Use fechas AAAA-MM-DD. Si no se han implementado medidas, describa esa situación y el avance real.")
    defaults = [legacy_action(report)] if report else []
    editor("Causas, medidas y compromisos", "acciones", ACCIONES, defaults)
    st.markdown("**Compromisos de reportes anteriores**")
    if previous:
        for r in previous:
            st.caption(f"{r.get('year', 2026)} · {r.get('periodo_label', r.get('reporte_id'))}: {r.get('compromisos', 'Sin descripción')}")
    f["avance_anteriores"] = st.text_area("Resumen de avance anterior o justificación de que no existen compromisos previos", value=f.get("avance_anteriores", ""), key=f"{key}_avance_anteriores", help=FIELD_HELP["avance_anteriores"])
    editor("Seguimiento individual de compromisos anteriores", "anteriores", ANTERIORES)
    if level == "rojo":
        st.subheader("Seguimiento y control de riesgo rojo")
        f["mecanismos"] = st.multiselect("Mecanismo competitivo identificado", ["Licitación", "Compra coordinada", "Convenio Marco", "Otro"], default=f.get("mecanismos", []), key=f"{key}_mecanismos", help=FIELD_HELP["mecanismos"])
        f["detalle_mecanismo"] = st.text_area("Detalle de regularización del abastecimiento", value=f.get("detalle_mecanismo", ""), key=f"{key}_detalle_mecanismo", help=FIELD_HELP["detalle_mecanismo"])
        f["fecha_reunion"] = text(st.date_input("Fecha programada de reunión técnica con SSMOCC", value=date_value(f.get("fecha_reunion")), format="DD-MM-YYYY", key=f"{key}_fecha_reunion", help=FIELD_HELP["fecha_reunion"]))
        st.caption("Monitoreo mensual hasta normalizar el indicador. Registre avances reales o controles programados; puede agregar meses adicionales.")
        defaults = [{"mes": f"Mes {i}", "fecha": "", "estado": "Programado", "observacion": ""} for i in range(1, 4)]
        editor("Controles mensuales", "mensual", MESES, defaults, "estado")
    else:
        st.subheader("Plan de seguimiento de riesgo amarillo")
        f["revision_recurrentes"] = choice("¿Se revisaron las compras recurrentes por TD?", ["Pendiente", "Sí", "No"], "revision_recurrentes")
        f["detalle_recurrentes"] = st.text_area("Ítems recurrentes, periodicidad y resultado de revisión o justificación", value=f.get("detalle_recurrentes", ""), key=f"{key}_detalle_recurrentes", help=FIELD_HELP["detalle_recurrentes"])
        f["seguimiento_reforzado"] = choice("¿Se compromete seguimiento reforzado?", ["Pendiente", "Sí", "No"], "seguimiento_reforzado")
        f["modalidad_seguimiento"] = st.text_input("Modalidad del seguimiento reforzado", value=f.get("modalidad_seguimiento", ""), key=f"{key}_modalidad_seguimiento", help=FIELD_HELP["modalidad_seguimiento"])
        f["fecha_seguimiento"] = text(st.date_input("Fecha del seguimiento reforzado", value=date_value(f.get("fecha_seguimiento")), format="DD-MM-YYYY", key=f"{key}_fecha_seguimiento", help=FIELD_HELP["fecha_seguimiento"]))
    st.subheader("Órdenes de compra y antecedentes de respaldo")
    f["por_informar"] = choice("Estado de OC con mecanismo Por Informar", ["Pendiente de revisión", "No existen", "Reclasificadas"], "por_informar")
    editor("Reclasificaciones según mecanismo real", "reclasificaciones", RECLASIFICACIONES)
    st.caption("Identifique las OC y los documentos de respaldo en un repositorio institucional con permisos adecuados. No registre contraseñas ni enlaces con credenciales.")
    editor("Índice de respaldos y detalle de OC", "respaldos", RESPALDOS)
    return f


def render_readonly(report):
    import streamlit as st
    f = report.get("formato_tipo", {})
    if not f:
        st.info("Reporte anterior conservado. Requiere completar los campos del Formato Tipo vigente.")
        return
    with st.expander("Plan de acción, seguimiento y respaldos registrados"):
        for field, columns, label in [("acciones", ACCIONES, "Plan de acción"), ("anteriores", ANTERIORES, "Compromisos anteriores"), ("mensual", MESES, "Control mensual"), ("reclasificaciones", RECLASIFICACIONES, "OC reclasificadas"), ("respaldos", RESPALDOS, "Respaldos")]:
            if f.get(field):
                st.markdown(f"**{label}**")
                st.dataframe(pd.DataFrame([{label: row.get(k, "") for k, label in columns.items()} for row in f[field]]), hide_index=True, use_container_width=True)
        for field, label in [("avance_anteriores", "Avance anterior"), ("fecha_reunion", "Reunión técnica"), ("detalle_mecanismo", "Regularización del abastecimiento"), ("detalle_recurrentes", "Compras recurrentes"), ("modalidad_seguimiento", "Seguimiento reforzado"), ("fecha_seguimiento", "Fecha de seguimiento"), ("por_informar", "OC Por Informar")]:
            if f.get(field):
                st.write(f"{label}: {f[field]}")


def format_tables(establishments, reports, year, rid, meta):
    """Incluye todos los riesgos obligatorios; nunca omite silenciosamente a los pendientes."""
    selected = {r["establecimiento_id"]: r for r in reports if int(r.get("year", 2026)) == int(year) and r.get("reporte_id") == rid}
    mandatory = {k: e for k, e in establishments.items() if e.get("nivel") in ("rojo", "amarillo")}
    controls = {r.get("establecimiento_id"): r for r in meta.get("pendientes", [])}
    indicators, comparisons, pending, warnings = [], [], [], []
    for eid, e in mandatory.items():
        r = selected.get(eid, {})
        included, official, same = revision(r, e)
        f = r.get("formato_tipo", {})
        indicators.append([e.get("codigo_deis", ""), e.get("rut", ""), e["nombre"], e.get("denominador", 0), e.get("numerador", 0), e.get("pct_2026", 0), e.get("pct_2025", 0), e.get("brecha", 0), e.get("variacion", 0), e["nivel"].capitalize()])
        comparisons.append([e["nombre"], "No incluido" if included == "No" else official, f.get("pct_minsal", "") if included == "Sí" else "—", e["nivel"].capitalize(), e.get("pct_2026", 0), e.get("variacion", 0), same, f.get("observacion_cotejo", "")])
        errors = completeness(r, e)
        if r.get("estado") != "enviado" or errors:
            control = controls.get(eid, {})
            status = "No recibido" if not r else "Borrador" if r.get("estado") != "enviado" else "Recibido; requiere completar antecedentes"
            pending.append([e["nombre"], status, fecha_chile(control.get("fecha")), control.get("observacion", "") or "; ".join(errors)])
            warnings.extend(f"{e['nombre']}: {err}" for err in errors)
    return mandatory, selected, indicators, comparisons, pending, warnings


def _set_paragraph(p, value):
    """Preserva el estilo del párrafo y del primer run de la plantilla."""
    if p.runs:
        p.runs[0].text = value
        for run in p.runs[1:]:
            run.text = ""
    else:
        p.add_run(value)


def _table_rows(table, rows):
    prototype = deepcopy(table.rows[1]._tr)
    for row in list(table.rows)[1:]:
        table._tbl.remove(row._tr)
    for values in rows or [["Sin antecedentes registrados"] + [""] * (len(table.columns) - 1)]:
        table._tbl.append(deepcopy(prototype))
        for cell, value in zip(table.rows[-1].cells, values):
            cell.text = fecha_chile(value) if isinstance(value, date) else text(value)


def generate_docx(establishments, reports, year, rid, period, meta, template=None):
    """Rellena y duplica los bloques de la plantilla oficial recibida, sin sustituir su estructura."""
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    doc = Document(template or Path(__file__).parent / "assets" / "formato_minsal_2026.docx")
    if meta.get("export_estado"):
        doc.paragraphs[5].insert_paragraph_before(meta["export_estado"])
    mandatory, selected, indicators, comparisons, pending, warnings = format_tables(establishments, reports, year, rid, meta)
    for p in doc.paragraphs:
        t = p.text
        if t.startswith("Servicio de Salud:"):
            _set_paragraph(p, "Servicio de Salud: Servicio de Salud Metropolitano Occidente")
        elif t.startswith("Reporte / período"):
            _set_paragraph(p, f"Reporte / período que se informa: {year} · {period['label']} · {period['periodo']}")
        elif t.startswith("Nombre de quien remite"):
            _set_paragraph(p, f"Nombre de quien remite / cargo: {meta.get('remitente', '')} / {meta.get('cargo', '')}")
        elif t.startswith("Correo electrónico:"):
            _set_paragraph(p, f"Correo electrónico: {meta.get('correo', '')}")
        elif t.startswith("Fecha de envío:"):
            _set_paragraph(p, f"Fecha de envío: {fecha_chile(meta.get('fecha_envio'))}")
        elif t.startswith("Visado por"):
            _set_paragraph(p, f"Visado por (Subdirector/a Administrativo, RRFF y Financieros): {meta.get('visado_por', '')}")
        elif t.startswith("Fecha:"):
            _set_paragraph(p, f"Fecha: {fecha_chile(meta.get('fecha_visado'))}")
        elif t.startswith("Plan de Monitoreo y Gestión"):
            _set_paragraph(p, f"Plan de Monitoreo y Gestión del Trato Directo — Año {year}")
    # Firma permanece disponible para el visado; nunca se fabrica una aprobación.
    original_tables = list(doc.tables)
    _table_rows(original_tables[0], [[number_chile(v, 0) if i in (3, 4) else number_chile(v) if i in (5, 6, 7, 8) else v for i, v in enumerate(row)] for row in indicators])
    _table_rows(original_tables[1], [[number_chile(v) if i in (2, 4, 5) and v not in ("", "—") else v for i, v in enumerate(row)] for row in comparisons])
    _table_rows(original_tables[7], pending)
    body = doc.element.body
    elements = list(body)
    starts = {}
    for i, element in enumerate(elements):
        if element.tag.endswith("}p"):
            t = Paragraph(element, doc).text
            for number in (4, 5, 6, 7):
                if t.startswith(f"{number}. "):
                    starts[number] = i
    prototypes = {"amarillo": [deepcopy(x) for x in elements[starts[4]+1:starts[5]]],
                  "rojo": [deepcopy(x) for x in elements[starts[5]+1:starts[6]]]}
    anchors = {number: elements[index] for number, index in starts.items()}
    for element in elements[starts[4]+1:starts[5]] + elements[starts[5]+1:starts[6]]:
        body.remove(element)

    def insert_block(eid, e, number, anchor):
        r = selected.get(eid, {})
        f = r.get("formato_tipo", {})
        block = [deepcopy(x) for x in prototypes[e["nivel"]]]
        tables = [Table(x, doc) for x in block if x.tag.endswith("}tbl")]
        included, _, _ = revision(r, e)
        identity = [e["nombre"], f"{e.get('codigo_deis', '')} / {e.get('rut', '')}", number_chile(e.get("pct_2026", 0)) + "%", number_chile(e.get("pct_2025", 0)) + "%", number_chile(e.get("variacion", 0)), e["nivel"].upper(), included]
        for row, value in zip(tables[0].rows, identity):
            row.cells[1].text = text(value)
        actions = f.get("acciones") or ([legacy_action(r)] if r else [])
        _table_rows(tables[1], [[i] + [fecha_chile(a.get(k)) if k == "plazo" else a.get(k, "") for k in ACCIONES] for i, a in enumerate(actions, 1)])
        if e["nivel"] == "rojo":
            _table_rows(tables[2], [[fecha_chile(a.get(k)) if k == "fecha" else a.get(k, "") for k in MESES] for a in f.get("mensual", [])])
        for x in block:
            if x.tag.endswith("}p"):
                p = Paragraph(x, doc)
                t = p.text
                if not t.strip() or t.startswith("Duplique este bloque"):
                    continue
                if t.startswith("Bloque tipo"):
                    _set_paragraph(p, f"Establecimiento N.º {number} — {e['nombre']} — {r.get('estado', 'No recibido')}")
                    p.paragraph_format.keep_with_next = True
                elif t.startswith("Aclaración sobre"):
                    _set_paragraph(p, "Aclaración sobre la clasificación: " + f.get("observacion_cotejo", ""))
                elif t.startswith("¿Se revisaron"):
                    _set_paragraph(p, "¿Se revisaron las compras recurrentes efectuadas mediante Trato Directo?: " + f.get("revision_recurrentes", "Pendiente"))
                elif t.startswith("Detalle de la revisión"):
                    _set_paragraph(p, "Detalle de la revisión: " + f.get("detalle_recurrentes", ""))
                elif t.startswith("Seguimiento reforzado comprometido"):
                    _set_paragraph(p, "Seguimiento reforzado comprometido: " + f.get("seguimiento_reforzado", "Pendiente"))
                elif t.startswith("Modalidad y fecha"):
                    _set_paragraph(p, f"Modalidad y fecha del seguimiento reforzado: {f.get('modalidad_seguimiento', '')} · {fecha_chile(f.get('fecha_seguimiento'))}")
                elif t.startswith("Mecanismo competitivo"):
                    _set_paragraph(p, "Mecanismo competitivo identificado: " + ", ".join(f.get("mecanismos", [])) + " — " + f.get("detalle_mecanismo", ""))
                elif t.startswith("Reunión de seguimiento"):
                    _set_paragraph(p, "Reunión de seguimiento técnico con el Servicio — fecha programada: " + fecha_chile(f.get("fecha_reunion")))
            anchor.addprevious(x)
        def paragraph(value):
            p = deepcopy(next(x for x in prototypes[e["nivel"]] if x.tag.endswith("}p")))
            paragraph = Paragraph(p, doc)
            paragraph.style = None
            props = p.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pPr")
            if props is not None:
                p.remove(props)
            _set_paragraph(paragraph, value)
            for run in paragraph.runs:
                run.bold = False
                from docx.shared import Pt
                run.font.size = Pt(9)
            paragraph.paragraph_format.space_after = Pt(6)
            anchor.addprevious(p)
        paragraph("Avance de compromisos de reportes anteriores: " + f.get("avance_anteriores", "Pendiente de informar"))
        for a in f.get("anteriores", []):
            paragraph(f"{a.get('periodo', '')} · {a.get('compromiso', '')} · Avance: {a.get('avance', '')} · Responsable: {a.get('responsable', '')} · Fecha: {fecha_chile(a.get('fecha'))}")
        paragraph("Estado OC Por Informar: " + f.get("por_informar", "Pendiente de revisión"))
        for a in f.get("reclasificaciones", []):
            paragraph(f"OC {a.get('oc', '')} · Mecanismo real: {a.get('mecanismo', '')} · Respaldo: {a.get('respaldo', '')}")
        paragraph("Antecedentes de respaldo y detalle de OC")
        for a in f.get("respaldos", []):
            paragraph(f"{a.get('oc', '')} · {a.get('descripcion', '')} · {a.get('url', '')}")
    n = 0
    counts = {4: 0, 5: 0, 6: 0}
    for eid, e in mandatory.items():
        included, _, _ = revision(selected.get(eid, {}), e)
        section = 6 if included == "No" else 4 if e["nivel"] == "amarillo" else 5
        n += 1
        counts[section] += 1
        insert_block(eid, e, n, anchors[section + 1])
    for section, count in counts.items():
        if not count:
            p = deepcopy(elements[6])
            _set_paragraph(Paragraph(p, doc), "Sin establecimientos que informar en esta sección.")
            anchors[section + 1].addprevious(p)
    # Conserva diseño, anchos y estilos originales y repite encabezados en tablas extensas.
    from docx.oxml import OxmlElement
    from docx.shared import Pt
    validation_section = False
    for p in doc.paragraphs:
        if p.text.startswith(("4. ", "5. ", "6. ")):
            p.paragraph_format.keep_with_next = True
        if p.text.startswith("8. "):
            validation_section = True
            p.paragraph_format.space_before = Pt(8)
        if validation_section:
            p.paragraph_format.space_after = Pt(3)
    for table in doc.tables:
        if len(table.columns) > 2:
            props = table.rows[0]._tr.get_or_add_trPr()
            props.append(OxmlElement("w:tblHeader"))
        if len(table.columns) == 2 or (len(table.columns) == 4 and len(table.rows) <= 5):
            # Mantiene juntas las fichas de identificación y los controles mensuales breves.
            for row in list(table.rows)[:-1]:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        p.paragraph_format.keep_with_next = True
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue(), warnings
